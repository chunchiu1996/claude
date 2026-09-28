import io

from openpyxl import Workbook, load_workbook

from shop.importer import COLUMNS, header_label, import_rows, read_upload, template_xlsx

from .conftest import product

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_template_has_products_instructions_and_examples():
    wb = load_workbook(io.BytesIO(template_xlsx()))
    assert wb.sheetnames == ["Products 产品", "Instructions 说明", "Example 示例"]
    products = wb["Products 产品"]
    headers = [c.value for c in products[1]]
    assert headers[0] == "SKU 货号" and "Quote only 仅询价" in headers and len(headers) == len(COLUMNS)
    assert products.max_row == 1 or all(c.value is None for c in products[2])  # blank sheet for the factory
    assert products["A1"].comment is not None and products.data_validations.dataValidation
    assert wb["Example 示例"]["A2"].value == "SPC-7MM-OAK"


def test_filled_in_template_imports(admin, db):
    """A factory fills in the template's Products sheet; the example sheet is ignored."""
    wb = load_workbook(io.BytesIO(template_xlsx()))
    ws = wb["Products 产品"]
    ws.append(["XL-TILE-1", "Wood-Look Porcelain 8x48 – Oak", "Tile", "Foshan New Factory", 31.2, 55, "box",
               "6 pcs / 15.5 sq ft", 1, 1200, "Ontario, CA", "50:28.5 | 200:26", "no", "Rectified plank tile.",
               "Size: 8x48 in; Finish: Matte", None, "yes", "yes"])
    ws.append(["XL-SLAB-2", "Quartz Slab – Pure White 2cm", "Countertops", "Foshan New Factory", None, None, "slab",
               None, 1, 30, None, None, "yes", None, None, None, None, None])
    out = io.BytesIO()
    wb.save(out)

    resp = admin.post("/admin/import", data={"file": (io.BytesIO(out.getvalue()), "shipment.xlsx", XLSX),
                                             "reference": "MSKU777"}, content_type="multipart/form-data")
    assert b"2 new, 0 updated" in resp.data and b"1,230 units received" in resp.data
    tile = product(db, "XL-TILE-1")
    assert (tile["price_cents"], tile["stock_qty"], tile["featured"], tile["quote_only"]) == (3120, 1200, 1, 0)
    assert db.execute("SELECT COUNT(*) FROM price_tiers WHERE product_id = ?", (tile["id"],)).fetchone()[0] == 2
    slab = product(db, "XL-SLAB-2")
    assert (slab["quote_only"], slab["price_cents"], slab["stock_qty"]) == (1, 0, 30)
    assert db.execute("SELECT 1 FROM products WHERE sku = 'SPC-7MM-OAK' AND stock_qty = 740").fetchone()  # example untouched


def test_factory_packing_list_with_title_rows_and_chinese_headers(db):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["佛山某某陶瓷有限公司 装箱单 Packing List"])
    ws.append(["Container: MSKU1234567"])
    ws.append([])
    ws.append(["货号", "品名", "单价", "数量", "类别"])
    ws.append([10023, "Glazed Tile 600x600", 2.8, 5000, "Tile"])  # numeric SKU cell from Excel
    out = io.BytesIO()
    wb.save(out)

    result = import_rows(db, read_upload("packing.xlsx", out.getvalue()))
    assert result.ok, result.errors
    tile = product(db, "10023")
    assert (tile["name"], tile["price_cents"], tile["stock_qty"]) == ("Glazed Tile 600x600", 280, 5000)


def test_row_errors_report_spreadsheet_row_numbers(db):
    rows = [["Packing list"], [header_label("sku"), header_label("name"), header_label("price")],
            ["NEW-A", "Has price", "10"], ["NEW-B", "Missing price", ""]]
    result = import_rows(db, rows)
    assert not result.ok and result.errors == [(4, "New SKU NEW-B needs a price (or set Quote only to yes)")]
    assert product(db, "NEW-A") is None  # all-or-nothing


def test_unreadable_files_give_friendly_errors():
    for name, raw in (("bad.xlsx", b"not a zip"), ("old.xls", b"\xd0\xcf\x11\xe0")):
        try:
            read_upload(name, raw)
        except ValueError as exc:
            assert ".xlsx" in str(exc)
        else:
            raise AssertionError("expected ValueError")


def test_admin_downloads_templates(admin):
    xlsx = admin.get("/admin/import/template.xlsx")
    assert xlsx.status_code == 200 and xlsx.mimetype == XLSX
    assert load_workbook(io.BytesIO(xlsx.data)).sheetnames[0] == "Products 产品"
    csv_text = admin.get("/admin/import/template.csv").data.decode()
    assert csv_text.startswith(",".join(COLUMNS))


def _fill_template(rows):
    """Type rows into the template's Products sheet the way a factory would in Excel."""
    wb = load_workbook(io.BytesIO(template_xlsx()))
    ws = wb["Products 产品"]
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row(**values):
    return [values.get(key) for key in COLUMNS]


def test_excel_auto_conversions_are_undone(db):
    import datetime

    raw = _fill_template([
        # Excel stores "50:2.49" as 00:50:02.49 and "100:10.80" as 01:40:10.80
        _row(sku="0001234", name="Oak Floor", price=3.5, stock=300, tiers=datetime.time(0, 50, 2, 490000)),
        _row(sku="TILE-1", name="Glazed Tile", price="$12.50", stock="1,200", tiers=datetime.time(1, 40, 10, 800000)),
    ])
    result = import_rows(db, read_upload("f.xlsx", raw))
    assert result.ok, result.errors
    assert product(db, "0001234")["price_cents"] == 350  # leading zeros kept
    tiers = db.execute("SELECT p.sku, t.min_qty, t.price_cents FROM price_tiers t JOIN products p ON p.id = t.product_id "
                       "WHERE p.sku IN ('0001234', 'TILE-1') ORDER BY p.sku").fetchall()
    assert [tuple(t) for t in tiers] == [("0001234", 50, 249), ("TILE-1", 100, 1080)]
    assert product(db, "TILE-1")["stock_qty"] == 1200


def test_sku_turned_into_a_date_is_explained(db):
    import datetime

    result = import_rows(db, read_upload("f.xlsx", _fill_template([_row(sku=datetime.datetime(2026, 3, 15), name="X", price=1)])))
    assert not result.ok and "looks like a date" in result.errors[0][1]


def test_chinese_categories_map_to_english_and_chinese_names_warn(db):
    raw = _fill_template([
        _row(sku="ZH-A", name="SPC Floor", category="地板", price=10),
        _row(sku="ZH-B", name="Wall Tile", category="瓷砖 / 石材", price=10),
        _row(sku="ZH-C", name="Faucet", category="Faucets & Fixtures 水龙头", price=10),
        _row(sku="ZH-D", name="石英石台面", category="台面", quote_only="是"),
    ])
    result = import_rows(db, read_upload("f.xlsx", raw))
    assert result.ok, result.errors
    cats = dict(db.execute("SELECT p.sku, c.name FROM products p JOIN categories c ON c.id = p.category_id "
                           "WHERE p.sku LIKE 'ZH-%'").fetchall())
    assert cats == {"ZH-A": "Flooring", "ZH-B": "Tile", "ZH-C": "Faucets & Fixtures", "ZH-D": "Countertops"}
    assert product(db, "ZH-D")["quote_only"] == 1
    assert [row for row, _ in result.warnings] == [5] and "Chinese" in result.warnings[0][1]


def test_template_keeps_risky_columns_as_text():
    ws = load_workbook(io.BytesIO(template_xlsx()))["Products 产品"]
    letters = {key: chr(ord("A") + i) for i, key in enumerate(COLUMNS)}
    for key in ("sku", "unit_note", "tiers"):
        assert ws.column_dimensions[letters[key]].number_format == "@", key
