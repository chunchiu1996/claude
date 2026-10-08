"""Bulk uploads (photos, big spreadsheets), paging, login protection, order emails and invoices."""
import io
import smtplib
import zipfile

from shop import notify, photos
from shop.importer import import_rows, read_upload

from .conftest import Client, image_bytes, product


def _upload(admin, *files):
    data = {"format": "json", "files": [(io.BytesIO(raw), name) for name, raw in files]}
    resp = admin.post("/admin/photos/upload", data=data, content_type="multipart/form-data",
                      headers={"Accept": "application/json"})
    assert resp.status_code == 200
    return resp.get_json()["results"]


# ---------------------------------------------------------------- photos


def test_photos_are_matched_to_products_by_file_name(admin, db, app):
    results = _upload(admin, ("spc-7mm-oak.JPG", image_bytes("JPEG", (3000, 2000))),
                      ("SPC-7MM-OAK-2.png", image_bytes()), ("SPC-7MM-OAK (3).png", image_bytes()),
                      ("NO-SUCH-SKU.jpg", image_bytes("JPEG")), ("TIL-POR-1224-CEM.jpg", b"not an image"))
    status = {r["file"]: r["status"] for r in results}
    assert status == {"spc-7mm-oak.JPG": "ok", "SPC-7MM-OAK-2.png": "ok", "SPC-7MM-OAK (3).png": "ok",
                      "NO-SUCH-SKU.jpg": "unmatched", "TIL-POR-1224-CEM.jpg": "error"}
    p = product(db, "SPC-7MM-OAK")
    assert p["image_url"].startswith("/media/") and p["image_url"].endswith(".jpg")
    assert [g["position"] for g in photos.gallery(db, p["id"])] == [2, 3]

    from PIL import Image

    folder = app.config["UPLOAD_FOLDER"]
    name = p["image_url"].rsplit("/", 1)[1]
    with Image.open(f"{folder}/{name}") as big, Image.open(f"{folder}/{name[:-4]}-sm.jpg") as small:
        assert max(big.size) == 1600 and max(small.size) == 600  # phone photos shrunk for the web

    # store: product card uses the small version; product page shows the gallery
    home = admin.get("/shop/flooring").data.decode()
    assert name[:-4] + "-sm.jpg" in home
    page = admin.get(f"/p/{p['slug']}").data.decode()
    assert page.count('class="gallery-thumb') == 3


def test_reuploading_a_photo_replaces_the_old_file(admin, db, app):
    _upload(admin, ("SPC-7MM-OAK.jpg", image_bytes("JPEG")))
    first = product(db, "SPC-7MM-OAK")["image_url"]
    _upload(admin, ("SPC-7MM-OAK.jpg", image_bytes("JPEG", color=(0, 0, 0))))
    second = product(db, "SPC-7MM-OAK")["image_url"]
    folder = app.config["UPLOAD_FOLDER"]
    assert first != second
    import os

    assert not os.path.exists(folder + "/" + first.rsplit("/", 1)[1])


def test_zip_of_photos(admin, db):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("photos/CAB-SHK-W-B30.jpg", image_bytes("JPEG"))
        z.writestr("photos/CAB-SHK-W-B30_2.jpg", image_bytes("JPEG"))
        z.writestr("__MACOSX/photos/._CAB-SHK-W-B30.jpg", b"junk")
        z.writestr("photos/readme.txt", b"hello")
    results = _upload(admin, ("photos.zip", buf.getvalue()))
    assert [(r["file"], r["status"], r["position"]) for r in results] == [
        ("CAB-SHK-W-B30.jpg", "ok", 1), ("CAB-SHK-W-B30_2.jpg", "ok", 2)]


def test_sku_ending_in_a_number_is_not_mistaken_for_a_gallery_photo(admin, db):
    admin.post("/admin/products/new", data={"sku": "VAN-2", "name": "Vanity two", "price": "10"})
    admin.post("/admin/products/new", data={"sku": "VAN", "name": "Vanity", "price": "10"})
    results = _upload(admin, ("VAN-2.jpg", image_bytes("JPEG")), ("VAN-3.jpg", image_bytes("JPEG")))
    assert [(r["sku"], r["position"]) for r in results] == [("VAN-2", 1), ("VAN", 3)]


def test_photos_page_lists_products_without_photos(admin, db):
    html = admin.get("/admin/photos").data.decode()
    assert "Products without a photo" in html and "SPC-7MM-OAK" in html
    _upload(admin, ("SPC-7MM-OAK.jpg", image_bytes("JPEG")))
    assert "<code>SPC-7MM-OAK</code>" not in admin.get("/admin/photos").data.decode()
    listing = admin.get("/admin/products?nophoto=1").data.decode()
    assert "SPC-7MM-OAK" not in listing and "SPC-7MM-GRY" in listing


def test_photo_upload_without_javascript_redirects_with_messages(admin):
    resp = admin.post("/admin/photos/upload", data={"files": [(io.BytesIO(image_bytes()), "nope.png")]},
                      content_type="multipart/form-data", follow_redirects=True)
    assert b"0 of 1 photos saved" in resp.data and b"nope.png" in resp.data


def test_product_form_gallery_add_and_remove(admin, db):
    p = product(db, "HDW-BARN-6FT")
    form = {"sku": p["sku"], "name": p["name"], "price": "79.00", "unit": "set", "moq": "1", "active": "1"}
    admin.post(f"/admin/products/{p['id']}", content_type="multipart/form-data", data={
        **form, "image": (io.BytesIO(image_bytes()), "main.png"),
        "extra_images": [(io.BytesIO(image_bytes()), "a.png"), (io.BytesIO(image_bytes()), "b.png")]})
    assert product(db, p["sku"])["image_url"]
    assert [g["position"] for g in photos.gallery(db, p["id"])] == [2, 3]
    admin.post(f"/admin/products/{p['id']}", data={**form, "remove_photo": "2"})
    assert [g["position"] for g in photos.gallery(db, p["id"])] == [3]


def test_photos_pasted_into_excel_are_attached(admin, db):
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage

    template = admin.get("/admin/import/template.xlsx").data
    wb = load_workbook(io.BytesIO(template))
    ws = wb[wb.sheetnames[0]]
    header_row = next(r for r in range(1, 10) if str(ws.cell(r, 1).value or "").lower().startswith("sku"))
    row = header_row + 1
    ws.cell(row, 1, "XL-PHOTO-1")
    ws.cell(row, 2, "Tile with pasted photo")
    ws.cell(row, 5, 12.5)
    for col in ("P", "Q"):
        picture = XLImage(io.BytesIO(image_bytes()))
        ws.add_image(picture, f"{col}{row}")
    out = io.BytesIO()
    wb.save(out)
    resp = admin.post("/admin/import", data={"file": (io.BytesIO(out.getvalue()), "factory.xlsx")},
                      content_type="multipart/form-data")
    html = resp.data.decode()
    assert "1 new" in html and "Photos from the spreadsheet: 2 of 2 saved" in html
    p = product(db, "XL-PHOTO-1")
    assert p["image_url"] and [g["position"] for g in photos.gallery(db, p["id"])] == [2]


def test_upload_too_big_gets_a_clear_message(app, admin):
    app.config["MAX_CONTENT_LENGTH"] = 1000
    resp = admin.post("/admin/photos/upload", data={"files": [(io.BytesIO(b"x" * 5000), "big.jpg")]},
                      content_type="multipart/form-data", headers={"Accept": "application/json"})
    assert resp.status_code == 413 and "too big" in resp.get_json()["error"]


# ---------------------------------------------------------------- big imports and paging


def _big_csv(n):
    lines = ["sku,name,category,price,quantity"]
    lines += [f"BULK-{i:05d},Bulk Tile {i},Tile,{10 + i % 50}.99,{i % 300}" for i in range(n)]
    return "\n".join(lines).encode()


def test_large_import_is_fast_and_pages_in_admin(admin, db):
    import time

    start = time.perf_counter()
    result = import_rows(db, read_upload("big.csv", _big_csv(3000)), reference="BIG-1", stock_mode="add")
    assert result.ok and result.created == 3000
    assert time.perf_counter() - start < 10

    html = admin.get("/admin/products?q=BULK").data.decode()
    assert html.count("<code>BULK-") == 50 and "Page 1 of 60" in html and "3,000 products" in html
    page_60 = admin.get("/admin/products?q=BULK&page=60").data.decode()
    assert page_60.count("<code>BULK-") == 50 and "Page 60 of 60" in page_60
    assert "Page 60 of 60" in admin.get("/admin/products?q=BULK&page=999").data.decode()  # clamped


def test_same_name_products_get_distinct_addresses(db):
    rows = read_upload("x.csv", b"sku,name,price\nA-1,Oak Plank,10\nB-2,Oak Plank,10\n")
    assert import_rows(db, rows, reference=None, stock_mode="add").ok
    slugs = {r["sku"]: r["slug"] for r in db.execute("SELECT sku, slug FROM products WHERE name = 'Oak Plank'")}
    assert slugs == {"A-1": "oak-plank", "B-2": "oak-plank-b-2"}


# ---------------------------------------------------------------- login protection


def test_login_locks_out_after_repeated_wrong_passwords(app):
    client = Client(app.test_client())
    for _ in range(5):
        assert client.post("/admin/login", data={"password": "guess"}).status_code == 200
    blocked = client.post("/admin/login", data={"password": "letmein"})  # even the right one, for a while
    assert blocked.status_code == 429 and b"Too many wrong passwords" in blocked.data
    other_ip = Client(app.test_client())
    resp = other_ip.post("/admin/login", data={"password": "letmein"}, environ_base={"REMOTE_ADDR": "10.0.0.9"})
    assert resp.status_code == 302


def test_successful_login_clears_failures(app):
    client = Client(app.test_client())
    for _ in range(4):
        client.post("/admin/login", data={"password": "guess"})
    assert client.post("/admin/login", data={"password": "letmein"}).status_code == 302
    for _ in range(4):
        client.post("/admin/login", data={"password": "guess"})
    assert client.post("/admin/login", data={"password": "letmein"}).status_code == 302


# ---------------------------------------------------------------- email, confirmations, invoices


class FakeSMTP:
    sent = []
    fail_login = False

    def __init__(self, host, port, timeout, **kwargs):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def starttls(self, context):
        pass

    def login(self, user, password):
        if FakeSMTP.fail_login:
            raise smtplib.SMTPAuthenticationError(535, b"bad credentials")

    def send_message(self, msg):
        FakeSMTP.sent.append(msg)


def _smtp(app, monkeypatch):
    FakeSMTP.sent, FakeSMTP.fail_login = [], False
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    app.config.update(SMTP_HOST="smtp.test", SMTP_USER="bot@shop.test", SMTP_PASSWORD="x",
                      SMTP_FROM="sales@liquidatorusa.com", NOTIFY_EMAIL="owner@shop.test",
                      PAYMENT_INSTRUCTIONS="Zelle: pay@liquidatorusa.com\\nCheck payable to Liquidator USA")


def _place_invoice_order(client, db):
    p = product(db, "HDW-BARN-6FT")
    client.post("/cart/add", data={"product_id": p["id"], "qty": 2})
    client.post("/checkout", data={"customer_name": "Pat Lee", "email": "pat@example.com", "phone": "555",
                                   "fulfillment": "pickup", "payment_method": "invoice"})
    return db.execute("SELECT * FROM orders WHERE email = 'pat@example.com'").fetchone()


def test_email_test_button(app, admin, monkeypatch):
    html = admin.get("/admin/email").data.decode()
    assert "Email is <strong>off</strong>" in html
    resp = admin.post("/admin/email", data={"to": "me@example.com"}, follow_redirects=True)
    assert b"Test email failed" in resp.data and b"SMTP_HOST is empty" in resp.data

    _smtp(app, monkeypatch)
    resp = admin.post("/admin/email", data={"to": "me@example.com"}, follow_redirects=True)
    assert b"Test email sent to me@example.com" in resp.data
    msg = FakeSMTP.sent[-1]
    assert msg["To"] == "me@example.com" and "Liquidator USA" in msg["From"] and "sales@liquidatorusa.com" in msg["From"]
    assert msg["Message-ID"] and msg["Date"]

    FakeSMTP.fail_login = True
    resp = admin.post("/admin/email", data={"to": "me@example.com"}, follow_redirects=True)
    assert b"rejected the login" in resp.data
    settings = admin.get("/admin/email").data.decode()
    assert "<code>SMTP_PASSWORD</code></td><td>set</td>" in settings  # the password itself is never shown


def test_order_confirmation_and_invoice(app, client, admin, db, monkeypatch):
    _smtp(app, monkeypatch)
    order = _place_invoice_order(client, db)
    to = [m["To"] for m in FakeSMTP.sent]
    assert to == ["owner@shop.test", "pat@example.com"]
    confirmation = FakeSMTP.sent[1].get_content()
    assert order["number"] in confirmation and "Total: $158.00" in confirmation and "/order/" in confirmation

    # customer can open a printable invoice from the order page
    order_page = client.get(f"/order/{order['number']}/{order['access_token']}").data.decode()
    invoice_link = f"/order/{order['number']}/{order['access_token']}/invoice"
    assert invoice_link in order_page
    invoice = client.get(invoice_link).data.decode()
    assert "INVOICE" in invoice and "Payment due" in invoice and "Zelle: pay@liquidatorusa.com" in invoice
    assert "$158.00" in invoice and "Check payable to Liquidator USA" in invoice
    assert client.get(f"/order/{order['number']}/wrong-token/invoice").status_code == 404

    # admin emails it
    resp = admin.post(f"/admin/orders/{order['id']}/email-invoice", follow_redirects=True)
    assert b"Invoice emailed to pat@example.com" in resp.data
    msg = FakeSMTP.sent[-1]
    assert msg["Subject"].startswith(f"Invoice {order['number']}") and msg["Reply-To"] == "owner@shop.test"
    text = msg.get_body(("plain",)).get_content()
    html = msg.get_body(("html",)).get_content()
    assert "How to pay:" in text and "Zelle" in text and "INVOICE" in html and "Print" not in html
    assert db.execute("SELECT invoice_sent_at FROM orders WHERE id = ?", (order["id"],)).fetchone()[0]

    admin.post(f"/admin/orders/{order['id']}", data={"status": "paid"})
    paid_invoice = admin.get(f"/admin/orders/{order['id']}/invoice").data.decode()
    assert "Total paid" in paid_invoice and "How to pay" not in paid_invoice


def test_invoice_email_failure_is_reported(app, client, admin, db, monkeypatch):
    order = _place_invoice_order(client, db)
    resp = admin.post(f"/admin/orders/{order['id']}/email-invoice", follow_redirects=True)
    assert b"Invoice not sent" in resp.data and b"Email isn" in resp.data


def test_card_payment_sends_confirmation_once(app, client, db, monkeypatch):
    from shop import payments

    sent = []
    monkeypatch.setattr(notify, "send_email", lambda to, subject, body, reply_to=None: sent.append((to, subject)))
    app.config["STRIPE_SECRET_KEY"] = "sk_test_x"
    monkeypatch.setattr(payments, "create_checkout_session", lambda *a, **k: ("cs_1", "https://stripe.test/pay"))
    p = product(db, "HDW-BARN-6FT")
    client.post("/cart/add", data={"product_id": p["id"], "qty": 1})
    client.post("/checkout", data={"customer_name": "Card Kim", "email": "kim@example.com", "phone": "555",
                                   "fulfillment": "pickup", "payment_method": "card"})
    assert sent == []  # nothing until Stripe says it's paid
    order = db.execute("SELECT * FROM orders WHERE email = 'kim@example.com'").fetchone()
    monkeypatch.setattr(payments, "retrieve_checkout_session",
                        lambda sid: {"payment_status": "paid", "client_reference_id": order["number"]})
    url = f"/order/{order['number']}/{order['access_token']}"
    client.get(url)
    client.get(url)
    assert [to for to, _ in sent] == ["sales@example.com", "kim@example.com"]
    assert "confirmed" in sent[1][1]


def test_notifications_are_sent_in_the_background_outside_tests(app, monkeypatch):
    import threading

    _smtp(app, monkeypatch)
    app.testing = False
    try:
        with app.test_request_context():
            assert notify.send_email("x@example.com", "Hi", "Body")
        for thread in threading.enumerate():
            if thread.name == "send-email":
                thread.join(5)
    finally:
        app.testing = True
    assert [m["To"] for m in FakeSMTP.sent] == ["x@example.com"]


def test_login_limit_uses_cloudflare_visitor_ip(app):
    app.config["CLIENT_IP_HEADER"] = "CF-Connecting-IP"
    client = Client(app.test_client())
    for _ in range(5):
        client.post("/admin/login", data={"password": "guess"}, headers={"CF-Connecting-IP": "203.0.113.5"})
    assert client.post("/admin/login", data={"password": "letmein"},
                       headers={"CF-Connecting-IP": "203.0.113.5"}).status_code == 429
    # another visitor coming through the same tunnel isn't locked out
    assert client.post("/admin/login", data={"password": "letmein"},
                       headers={"CF-Connecting-IP": "198.51.100.7"}).status_code == 302


def test_iphone_heic_photos(admin, db):
    import pytest

    pillow_heif = pytest.importorskip("pillow_heif")
    from PIL import Image

    pillow_heif.register_heif_opener()
    out = io.BytesIO()
    try:
        Image.new("RGB", (64, 48), (200, 120, 40)).save(out, "HEIF")
    except (KeyError, OSError, ValueError):
        pytest.skip("this pillow-heif build can't write HEIC")
    results = _upload(admin, ("LGT-VAN-3L-BLK.HEIC", out.getvalue()))
    assert results[0]["status"] == "ok" and product(db, "LGT-VAN-3L-BLK")["image_url"].endswith(".jpg")


def test_broken_zip_entries_and_truncated_photos_are_reported(admin, db):
    good = image_bytes("JPEG", (400, 300))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("SPC-7MM-OAK.jpg", good)
        z.writestr("SPC-7MM-GRY.jpg", good[: len(good) // 3])  # truncated download
        z.writestr("LAM-12MM-HIC.jpg", good)
    raw = bytearray(buf.getvalue())
    # corrupt the compressed bytes of the last entry so unzipping it fails its checksum
    info = zipfile.ZipFile(io.BytesIO(bytes(raw))).getinfo("LAM-12MM-HIC.jpg")
    data_start = info.header_offset + 30 + len(info.filename) + len(info.extra)
    for i in range(data_start + 5, data_start + 25):
        raw[i] ^= 0xFF
    results = _upload(admin, ("photos.zip", bytes(raw)))
    status = {r["file"]: r["status"] for r in results}
    assert status == {"SPC-7MM-OAK.jpg": "ok", "SPC-7MM-GRY.jpg": "error", "LAM-12MM-HIC.jpg": "error"}
