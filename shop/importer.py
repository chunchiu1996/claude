"""Bulk product/stock import from the spreadsheets factories send with each shipment (.xlsx or .csv).

The import is all-or-nothing: if any row has an error nothing is saved, so re-uploading a
corrected file never double-counts received stock.
"""
import csv
import datetime
import io
import re

from .catalog import (
    get_or_create_category,
    get_or_create_supplier,
    normalize_specs,
    record_movement,
    set_tiers,
    slugify,
    unique_slug,
)
from .pricing import parse_money, parse_tiers

REQUIRED, NEW_ONLY, RECOMMENDED, OPTIONAL = "Always 必填", "New products 新产品必填", "Recommended 建议填写", "Optional 选填"

# key, English header, Chinese header, required, help (EN), help (ZH), example
COLUMN_SPECS = [
    ("sku", "SKU", "货号", REQUIRED,
     "Your unique product code. Rows are matched by SKU: new SKUs are created, existing SKUs are updated.",
     "产品唯一编码。按货号匹配：新货号新建产品，已有货号更新产品。", "SPC-7MM-OAK"),
    ("name", "Product name", "产品名称", NEW_ONLY,
     "Name shown to customers, in English.", "英文产品名称（顾客可见）。",
     "Rigid Core SPC Vinyl Plank - Natural Oak, 7mm"),
    ("category", "Category", "类别", RECOMMENDED,
     "e.g. Flooring, Tile, Kitchen Cabinets, Bathroom Vanities, Faucets & Fixtures, Lighting, Doors & Hardware. "
     "A new name creates a new category.",
     "例如 Flooring 地板、Tile 瓷砖、Kitchen Cabinets 橱柜、Lighting 灯具。填写新名称会自动新建类别。", "Flooring"),
    ("supplier", "Factory", "工厂名称", RECOMMENDED,
     "The factory that owns this consignment stock.", "寄售库存所属工厂。", "Changzhou Example Flooring Co."),
    ("price", "Price USD", "零售价 美元", NEW_ONLY,
     "Selling price per unit in US dollars. Can be left blank when Quote only = yes.",
     "每单位美元售价。如“仅询价”填 yes，可留空。", "42.50"),
    ("compare_at_price", "Typical retail USD", "市场参考价 美元", OPTIONAL,
     "Typical US store price, shown crossed out to highlight the saving.", "美国零售店参考价，用于显示折扣。", "66.00"),
    ("unit", "Unit", "销售单位", OPTIONAL,
     "How it is sold: box, piece, set, sheet, carton, pallet, roll… (default: piece).",
     "销售单位：box 箱、piece 件、set 套、sheet 片、pallet 托……（默认 piece）。", "box"),
    ("unit_note", "Unit details", "单位说明", OPTIONAL,
     "What one unit contains.", "每单位包含的内容。", "23.64 sq ft per box"),
    ("moq", "Min order qty", "最小起订量", OPTIONAL,
     "Minimum units per order (default 1).", "每单最少数量（默认 1）。", "1"),
    ("stock", "Quantity", "数量", OPTIONAL,
     "Units arriving in this shipment (or total on hand when doing a stock count).",
     "本批到货数量（盘点时填写总库存）。", "860"),
    ("warehouse", "US warehouse", "美国仓库", OPTIONAL,
     "City/state of the US warehouse holding the stock.", "库存所在美国仓库（城市/州）。", "Ontario, CA"),
    ("tiers", "Volume pricing", "阶梯价格", OPTIONAL,
     "QTY:PRICE pairs separated by |  (buy 20+ for $38.99, 100+ for $35.45).",
     "数量:单价，多个用 | 分隔（如 20 件以上 38.99）。", "20:38.99 | 100:35.45"),
    ("quote_only", "Quote only", "仅询价", OPTIONAL,
     "yes = hide the price; customers send a quote request instead.", "yes = 不显示价格，顾客提交询价。", "no"),
    ("description", "Description", "产品描述", OPTIONAL,
     "A few sentences for customers, in English.", "英文产品描述。",
     "Waterproof click-lock plank with attached IXPE pad."),
    ("specs", "Specifications", "规格参数", OPTIONAL,
     "Key: Value pairs separated by ;", "参数名: 参数值，多个用 ; 分隔。",
     "Thickness: 7mm; Wear layer: 20 mil; Plank size: 7.2 x 60 in"),
    ("image_url", "Image URL", "图片链接", OPTIONAL,
     "Link to a product photo. You can also upload photos later in admin.", "产品图片链接（也可以之后在后台上传）。",
     "https://example.com/oak.jpg"),
    ("featured", "Featured", "首页推荐", OPTIONAL, "yes = show on the home page.", "yes = 在首页展示。", "no"),
    ("active", "Visible", "上架", OPTIONAL,
     "no = hide from the shop (new products are visible by default).", "no = 不在网店显示（新产品默认上架）。", "yes"),
]
COLUMNS = [spec[0] for spec in COLUMN_SPECS]

EXAMPLE_ROWS = [
    {"sku": "SPC-7MM-OAK", "name": "Rigid Core SPC Vinyl Plank - Natural Oak, 7mm", "category": "Flooring",
     "supplier": "Changzhou Example Flooring Co.", "price": 42.5, "compare_at_price": 66, "unit": "box",
     "unit_note": "23.64 sq ft per box", "moq": 1, "stock": 860, "warehouse": "Ontario, CA",
     "tiers": "20:38.99 | 100:35.45", "quote_only": "no",
     "description": "Waterproof click-lock plank with attached IXPE pad.",
     "specs": "Thickness: 7mm; Wear layer: 20 mil; Plank size: 7.2 x 60 in", "featured": "yes", "active": "yes"},
    {"sku": "QTZ-SLAB-CAL-3CM", "name": "Quartz Slab 126x63 in - Calacatta Gold, 3cm", "category": "Countertops",
     "supplier": "Foshan Example Stone Co.", "unit": "slab", "unit_note": "126 x 63 in, polished", "moq": 1,
     "stock": 45, "warehouse": "Ontario, CA", "quote_only": "yes",
     "description": "Engineered quartz slab for kitchen countertops and islands.",
     "specs": "Size: 126 x 63 in; Thickness: 3cm; Finish: Polished", "active": "yes"},
    {"sku": "FAU-KIT-PD-BN", "name": "Pull-Down Kitchen Faucet - Brushed Nickel", "category": "Faucets & Fixtures",
     "supplier": "Kaiping Example Sanitary Ware", "price": 68, "compare_at_price": 149, "unit": "piece", "moq": 1,
     "stock": 520, "warehouse": "Ontario, CA", "tiers": "10:61 | 50:55", "quote_only": "no",
     "specs": "Material: Solid brass; Flow rate: 1.8 GPM; Certification: cUPC", "active": "yes"},
]

_EXTRA_ALIASES = {
    "retail_price": "price", "unit_price": "price", "qty": "stock", "quantity": "stock", "factory": "supplier",
    "min_order": "moq", "volume_pricing": "tiers", "image": "image_url", "price_on_request": "quote_only",
    "型号": "sku", "编号": "sku", "品名": "name", "价格": "price", "单价": "price", "库存": "stock",
    "到货数量": "stock", "工厂": "supplier", "供应商": "supplier", "规格": "specs", "描述": "description",
}
TRUE_VALUES = {"1", "y", "yes", "true", "x", "是", "✓"}

# Factories often write categories in Chinese; the English store needs English category names.
CHINESE_CATEGORIES = {
    "地板": "Flooring", "木地板": "Flooring", "石塑地板": "Flooring", "spc地板": "Flooring", "强化地板": "Flooring",
    "复合地板": "Flooring", "瓷砖": "Tile", "地砖": "Tile", "墙砖": "Tile", "石材": "Tile", "马赛克": "Tile",
    "橱柜": "Kitchen Cabinets", "浴室柜": "Bathroom Vanities", "水龙头": "Faucets & Fixtures",
    "龙头": "Faucets & Fixtures", "卫浴": "Faucets & Fixtures", "花洒": "Faucets & Fixtures", "灯具": "Lighting",
    "照明": "Lighting", "灯": "Lighting", "门": "Doors & Hardware", "五金": "Doors & Hardware",
    "门锁": "Doors & Hardware", "台面": "Countertops", "石英石": "Countertops", "岩板": "Countertops",
}
_CJK = re.compile(r"[\u3400-\u9fff]")


def english_category(value):
    """'瓷砖 / 石材' -> 'Tile'; 'Flooring 地板' -> 'Flooring'; English names are kept as they are."""
    english = _CJK.sub("", value).strip(" /、,，()（）")
    if english and english != value.strip():
        return english
    for part in re.split(r"[/、,，\s]+", value.lower()):
        part = part.strip("()（）")
        if part in CHINESE_CATEGORIES:
            return CHINESE_CATEGORIES[part]
    return value


def _squash(text):
    return re.sub(r"[\s_\-()（）/:：.*#]+", "", str(text or "").lower())


HEADER_MAP = {}
for _key, _en, _zh, *_ in COLUMN_SPECS:
    for _variant in (_key, _en, _zh, f"{_en} {_zh}", f"{_zh} {_en}"):
        HEADER_MAP[_squash(_variant)] = _key
HEADER_MAP.update({_squash(alias): key for alias, key in _EXTRA_ALIASES.items()})


def header_label(key):
    spec = next(s for s in COLUMN_SPECS if s[0] == key)
    return f"{spec[1]} {spec[2]}"


def _norm_header(name):
    squashed = _squash(name)
    return HEADER_MAP.get(squashed, squashed)


class ImportResult:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.units_received = 0
        self.errors = []  # (spreadsheet row number, message) — any error means nothing is imported
        self.warnings = []  # (row number, message) — imported, but worth a look

    @property
    def ok(self):
        return not self.errors


def decode(raw):
    """Excel on Chinese-locale Windows often saves CSV as GBK, so fall back to GB18030."""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (datetime.time, datetime.timedelta)):
        # Excel turns a volume price typed as "50:2.49" into the time 00:50:02.49 — turn it back.
        if isinstance(value, datetime.time):
            minutes, seconds = value.hour * 60 + value.minute, value.second + value.microsecond / 1e6
        else:
            minutes, seconds = divmod(value.total_seconds(), 60)
        return f"{int(minutes)}:{seconds:.2f}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_upload(filename, raw):
    """Turn an uploaded .xlsx or .csv into a list of rows (lists of strings)."""
    if (filename or "").lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many different types for corrupt files
            raise ValueError("Couldn't read that Excel file. Save it as .xlsx (or CSV) and try again.") from exc
        sheet = next(
            (ws for ws in workbook.worksheets if _squash(ws.title).startswith("products") or "产品" in ws.title),
            workbook.worksheets[0],
        )
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    if (filename or "").lower().endswith(".xls"):
        raise ValueError("Old .xls files aren't supported. In Excel choose File → Save As → .xlsx, then upload again.")
    return list(csv.reader(io.StringIO(decode(raw))))


def import_products(db, text, reference=None, stock_mode="add"):
    """Import CSV text (see import_rows)."""
    return import_rows(db, list(csv.reader(io.StringIO(text))), reference, stock_mode)


def import_rows(db, rows, reference=None, stock_mode="add"):
    """stock_mode 'add': the quantity column is units just received (a new shipment).
    stock_mode 'set': the quantity column is the full count on hand (a stocktake)."""
    result = ImportResult()
    # Factory packing lists often have title rows above the real header, so look for it.
    header_index = next(
        (i for i, row in enumerate(rows[:15]) if "sku" in {_norm_header(c) for c in row}), None
    )
    if header_index is None:
        result.errors.append((1, "Couldn't find a 'SKU 货号' column. Download the template to see the expected columns."))
        return result
    headers = [_norm_header(c) for c in rows[header_index]]

    seen = set()
    try:
        for rownum, raw_row in enumerate(rows[header_index + 1:], start=header_index + 2):
            row = {}
            for key, value in zip(headers, raw_row, strict=False):  # rows may be shorter/longer than the header
                if key in COLUMNS and key not in row:
                    row[key] = _cell(value)
            if not any(row.values()):
                continue
            try:
                _import_row(db, row, reference, stock_mode, result, seen, rownum)
            except ValueError as exc:
                result.errors.append((rownum, str(exc)))
        if result.errors:
            db.rollback()
        else:
            db.commit()
    except Exception:
        db.rollback()
        raise
    return result


def _int(value, field):
    try:
        return int(str(value).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"{field} must be a whole number, got {value!r}") from None


def _import_row(db, row, reference, stock_mode, result, seen, rownum):
    sku = row.get("sku", "")
    if not sku:
        raise ValueError("SKU is required")
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}( 00:00:00)?", sku):
        raise ValueError(f"SKU {sku} looks like a date: Excel converted it. Format the SKU column as Text and retype it")
    if sku.lower() in seen:
        raise ValueError(f"SKU {sku} appears more than once in this file")
    seen.add(sku.lower())

    existing = db.execute("SELECT * FROM products WHERE sku = ? COLLATE NOCASE", (sku,)).fetchone()
    fields = {}
    if row.get("name"):
        fields["name"] = row["name"]
        if _CJK.search(row["name"]):
            result.warnings.append((rownum, f"{sku}: the product name is in Chinese. The English store will show it "
                                            "as written, so add an English name when you can"))
    if row.get("price"):
        fields["price_cents"] = parse_money(row["price"])
    if row.get("compare_at_price"):
        fields["compare_at_cents"] = parse_money(row["compare_at_price"])
    for col in ("unit", "unit_note", "warehouse", "description", "image_url"):
        if row.get(col):
            fields[col] = row[col]
    if row.get("specs"):
        fields["specs"] = normalize_specs(row["specs"])
    if row.get("moq"):
        fields["moq"] = _int(row["moq"], "Min order qty")
        if fields["moq"] < 1:
            raise ValueError("Min order qty must be at least 1")
    for flag in ("featured", "active", "quote_only"):
        if row.get(flag):
            fields[flag] = 1 if row[flag].lower() in TRUE_VALUES else 0
    if row.get("category"):
        category = english_category(row["category"])
        if _CJK.search(category):
            result.warnings.append((rownum, f"{sku}: category “{category}” is in Chinese. Rename it under "
                                            "Admin → Categories, or use an English category name"))
        fields["category_id"] = get_or_create_category(db, category)
    if row.get("supplier"):
        fields["supplier_id"] = get_or_create_supplier(db, row["supplier"])
    tiers = parse_tiers(row["tiers"]) if row.get("tiers") else None
    stock = _int(row["stock"], "Quantity") if row.get("stock") else None
    if stock is not None and stock < 0:
        raise ValueError("Quantity can't be negative")

    if existing is None:
        if "name" not in fields:
            raise ValueError(f"New SKU {sku} needs a product name")
        if "price_cents" not in fields:
            if not fields.get("quote_only"):
                raise ValueError(f"New SKU {sku} needs a price (or set Quote only to yes)")
            fields["price_cents"] = 0
        fields["sku"] = sku
        fields["slug"] = unique_slug(db, "products", slugify(fields["name"]))
        product_id = db.execute(
            f"INSERT INTO products ({', '.join(fields)}) VALUES ({', '.join('?' for _ in fields)})",
            list(fields.values()),
        ).lastrowid
        current_stock = 0
        result.created += 1
    else:
        product_id = existing["id"]
        current_stock = existing["stock_qty"]
        if fields:
            assignments = ", ".join(f"{col} = ?" for col in fields)
            db.execute(
                f"UPDATE products SET {assignments}, updated_at = datetime('now') WHERE id = ?",
                [*fields.values(), product_id],
            )
        result.updated += 1

    final = db.execute("SELECT price_cents, quote_only FROM products WHERE id = ?", (product_id,)).fetchone()
    if not final["quote_only"] and final["price_cents"] <= 0:
        raise ValueError(f"SKU {sku} has no price. Add one, or set Quote only to yes")

    if tiers is not None:
        set_tiers(db, product_id, tiers)

    if stock is not None:
        if stock_mode == "set":
            delta = stock - current_stock
            if delta:
                record_movement(db, product_id, delta, "adjust", reference, "Stock count import")
        elif stock:
            record_movement(db, product_id, stock, "receive", reference, "Spreadsheet import")
            result.units_received += stock


# ---------------------------------------------------------------- templates


def template_csv():
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    for example in EXAMPLE_ROWS:
        writer.writerow([_cell(example.get(key)) for key in COLUMNS])
    return out.getvalue()


def template_xlsx():
    """A bilingual Excel template to send to factories: a blank Products sheet, instructions, and examples."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    brown, orange = PatternFill("solid", fgColor="6B3A22"), PatternFill("solid", fgColor="D9681E")
    white_bold, bold = Font(bold=True, color="FFFFFF"), Font(bold=True)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))
    last_row = 2000

    def header_row(ws):
        for col, (key, en, zh, required, help_en, help_zh, example) in enumerate(COLUMN_SPECS, start=1):
            cell = ws.cell(row=1, column=col, value=f"{en} {zh}")
            cell.font = white_bold
            cell.fill = orange if required in (REQUIRED, NEW_ONLY) else brown
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.comment = Comment(f"{help_en}\n{help_zh}\nExample 示例: {example}", "Template", width=320, height=140)
            width = {"name": 42, "description": 44, "specs": 44, "tiers": 22, "unit_note": 22, "supplier": 28}
            ws.column_dimensions[get_column_letter(col)].width = width.get(key, 16)
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "B2"

    wb = Workbook()
    products = wb.active
    products.title = "Products 产品"
    header_row(products)
    col = {key: get_column_letter(i) for i, key in enumerate(COLUMNS, start=1)}
    # Text columns: stop Excel turning SKU 00123 into 123, "3-15" into a date, "50:2.49" into a time.
    for key in ("sku", "unit_note", "tiers"):
        products.column_dimensions[col[key]].number_format = "@"

    def validation(dv, *keys):
        products.add_data_validation(dv)
        for key in keys:
            dv.add(f"{col[key]}2:{col[key]}{last_row}")

    validation(DataValidation(type="list", formula1='"yes,no"', allow_blank=True), "quote_only", "featured", "active")
    validation(DataValidation(type="list", formula1='"box,piece,set,sheet,carton,pallet,roll,pair,slab,sq ft"',
                              allow_blank=True, showErrorMessage=False), "unit")
    validation(DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Price 价格",
                              error="Enter a number in US dollars, e.g. 42.50\n请输入美元数字，如 42.50"),
               "price", "compare_at_price")
    validation(DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Quantity 数量",
                              error="Enter a whole number, e.g. 860\n请输入整数，如 860"), "stock", "moq")

    guide = wb.create_sheet("Instructions 说明")
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 22
    guide.column_dimensions["C"].width = 60
    guide.column_dimensions["D"].width = 44
    guide.column_dimensions["E"].width = 34
    notes = [
        ("How to fill in this template", "填写说明"),
        ("1. Put one product per row on the “Products 产品” sheet. Don't change the header row.",
         "1. 在“Products 产品”表中每行填写一个产品，请勿修改第一行表头。"),
        ("2. Orange columns are required for new products. Brown columns are optional.",
         "2. 橙色列为新产品必填，棕色列为选填。"),
        ("3. Prices are in US dollars, numbers only (no $ sign).", "3. 价格为美元，只填数字（不要 $ 符号）。"),
        ("4. Quantity = units arriving in this shipment. Tell us the container number when you send the file.",
         "4. 数量 = 本批到货数量。发送文件时请注明柜号。"),
        ("5. Set Quote only = yes to hide the price and let customers ask for a quote.",
         "5. “仅询价”填 yes：网店不显示价格，顾客提交询价。"),
        ("6. Hover over a header cell to see help. See the “Example 示例” sheet for filled-in rows.",
         "6. 鼠标悬停表头可查看说明；“Example 示例”表中有填写范例。"),
        ("7. Save as .xlsx and send it back to us.", "7. 保存为 .xlsx 文件发回给我们。"),
    ]
    for r, (en, zh) in enumerate(notes, start=1):
        guide.cell(row=r, column=1, value=en).font = bold if r == 1 else Font()
        guide.cell(row=r, column=3, value=zh).font = bold if r == 1 else Font()
    start = len(notes) + 2
    for c, title in enumerate(["Column 列名", "Required 是否必填", "Description", "说明", "Example 示例"], start=1):
        cell = guide.cell(row=start, column=c, value=title)
        cell.font, cell.fill = white_bold, brown
    for r, (_key, en, zh, required, help_en, help_zh, example) in enumerate(COLUMN_SPECS, start=start + 1):
        for c, value in enumerate([f"{en} {zh}", required, help_en, help_zh, example], start=1):
            cell = guide.cell(row=r, column=c, value=value)
            cell.alignment, cell.border = wrap_top, thin
        if required in (REQUIRED, NEW_ONLY):
            guide.cell(row=r, column=2).font = Font(bold=True, color="D9681E")

    examples = wb.create_sheet("Example 示例")
    header_row(examples)
    for r, example in enumerate(EXAMPLE_ROWS, start=2):
        for c, key in enumerate(COLUMNS, start=1):
            examples.cell(row=r, column=c, value=example.get(key))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
